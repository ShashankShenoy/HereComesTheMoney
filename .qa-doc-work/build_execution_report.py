from collections import Counter
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree
import json

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from test_cases import GROUPS

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / 'Downloads/integrated'
EVIDENCE = BASE / '.runtime/oracle-m12'
OUT = BASE / 'docs/MODULE-1-AND-2-ORACLE-TEST-EXECUTION.docx'
SUITE = ElementTree.parse(EVIDENCE / 'backend-final.xml').getroot()
PREFLIGHT = json.loads((EVIDENCE / 'schema-preflight.json').read_text())
CASES = [case for group in GROUPS for case in group['cases']]

# A case is complete only when every step and variant in the separate test plan
# has evidence. Focused automated checks do not silently complete a broad case.
NOT_RUN = {
    'M1-OVR-02': 'Fresh, unbootstrapped Oracle schema unavailable; repeat setup was rejected as expected.',
    'M1-AUTH-09': 'Contaminated signup-role configuration requires a disposable isolated Oracle schema.',
    'M2-LIFE-04': 'No synthetic account with every financial blocker was prepared; shared balances were protected.',
    'X-FLOW-06': 'A controlled identity/gateway outage and recovery was not run.',
}

PARTIAL_DETAILS = {
    'M1-ACC-08': 'Oracle overlap rejected and adjacent validity windows accepted; simultaneous decisions remain open.',
    'X-FLOW-03': 'Oracle stale edit/review requests denied after role retirement; open-browser dialog and restore flow remain open.',
}

BROWSER = {
    'M1-AUTH', 'M1-OVR', 'M1-USR', 'M1-ROLE', 'M1-ACC', 'M1-SEC',
    'M1-MENU', 'M1-AUD', 'M1-KEY', 'M2-DIR', 'M2-PRO', 'M2-KYC',
    'M2-DOC', 'M2-REL', 'M2-AUD', 'X-UI',
}
GROUP_EVIDENCE = {
    'M1-AUTH': 'Oracle authentication, signup and MFA methods; sign-in UI',
    'M1-OVR': 'Oracle dashboard/audit method; Overview UI',
    'M1-USR': 'Oracle identity, status, session and paging methods; Users UI',
    'M1-ROLE': 'Oracle role, permission and authority methods; Roles UI',
    'M1-ACC': 'Oracle access decision, scope and paging methods; queue UI',
    'M1-SEC': 'Oracle session, MFA and authorization methods; My security UI',
    'M1-MENU': 'Oracle menu contract method; menu catalogue UI',
    'M1-AUD': 'Oracle audit method; security and financial audit UI',
    'M1-KEY': 'Oracle key-rotation method; key form and clear UI',
    'M2-DIR': 'Oracle customer directory/scope method; directory and tab UI',
    'M2-PRO': 'Oracle profile, facts and boundary methods; Profile tab UI',
    'M2-KYC': 'Oracle completeness, approval, rework and expiry methods; KYC tab UI',
    'M2-DOC': 'Oracle upload, review, integrity and download methods; upload UI',
    'M2-LIFE': 'Oracle lifecycle and closure methods',
    'M2-REL': 'Oracle consent and relationship method; Relationships tab UI',
    'M2-AUD': 'Oracle audit/scope methods; customer Audit tab UI',
    'X-FLOW': 'Related Oracle integration methods; full multi-session variant open',
    'X-UI': 'Oracle browser checks and 56 Node checks; variant coverage open',
}

def group_key(case_id):
    return '-'.join(case_id.split('-')[:-1])

doc = Document()
section = doc.sections[0]
section.page_width, section.page_height = Inches(8.27), Inches(11.69)
section.top_margin = section.bottom_margin = Inches(.65)
section.left_margin = section.right_margin = Inches(.7)
section.header_distance = section.footer_distance = Inches(.3)
for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2'):
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Arial')
doc.styles['Normal'].font.size = Pt(9)
doc.styles['Normal'].paragraph_format.space_after = Pt(5)
doc.styles['Title'].font.size = Pt(24)
doc.styles['Title'].font.bold = True
doc.styles['Heading 1'].font.size = Pt(15)
doc.styles['Heading 2'].font.size = Pt(11)
doc.core_properties.title = 'Module 1 and 2 Oracle Test Execution'
doc.core_properties.subject = 'Oracle-backed IAM and CIF/KYC execution evidence and gaps'
doc.core_properties.author = 'Moneybags QA'

def paragraph(text='', style=None):
    p = doc.add_paragraph(style=style)
    p.add_run(text)
    return p

def heading(text, level=1):
    return doc.add_heading(text, level)

def table(headers, rows, widths):
    t = doc.add_table(rows=1, cols=len(headers))
    t.autofit = False
    for i, width in enumerate(widths):
        t.columns[i].width = Inches(width)
        t.rows[0].cells[i].text = headers[i]
    for values in rows:
        cells = t.add_row().cells
        for i, value in enumerate(values):
            cells[i].text = str(value)
    for index, row in enumerate(t.rows):
        trp = row._tr.get_or_add_trPr()
        trp.append(OxmlElement('w:cantSplit'))
        if index == 0:
            trp.append(OxmlElement('w:tblHeader'))
        for i, cell in enumerate(row.cells):
            cell.width = Inches(widths[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            sh = OxmlElement('w:shd')
            sh.set(qn('w:fill'), 'DCE6EF' if index == 0 else ('F5F7F9' if index % 2 == 0 else 'FFFFFF'))
            cell._tc.get_or_add_tcPr().append(sh)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(1)
                for run in p.runs:
                    run.font.size = Pt(8)
                    run.bold = index == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return t

footer = section.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Moneybags QA | Oracle run | 14 October 2026').font.size = Pt(8)

paragraph('Module 1 and 2 Oracle Test Execution', 'Title')
paragraph('Run 14 October 2026 | IAM and Customers & KYC', 'Subtitle')
paragraph('40 of 40 Oracle acceptance methods passed; 56 of 56 database-free frontend checks passed. The 118-case test plan remains an execution worksheet: focused evidence does not close a case whose other steps or variants have not been run.')
table(['Check', 'Result', 'Evidence'], [
    ('Database', 'Oracle metadata asserted in the acceptance suite; no demo H2 tests run', 'backend-final.xml'),
    ('Oracle acceptance', f"{SUITE.attrib['tests']} passed; {SUITE.attrib['failures']} failed; {SUITE.attrib['errors']} errors; {SUITE.attrib['skipped']} skipped", 'backend-final.xml'),
    ('Frontend and edge', '56 passed; 0 failed', 'frontend-node-tests.log'),
    ('Oracle schema preflight', f"{sum(r['status']=='PASS' for r in PREFLIGHT['results'])} pass; {sum(r['status']=='FAIL' for r in PREFLIGHT['results'])} fail", 'schema-preflight.json'),
], [1.55, 3.72, 1.6])
paragraph('Environment: Oracle-backed gateway and services on local QA ports 18090–18097; Oracle JET frontend on 15173. Customer signup was disabled in the live UI and enabled only inside rollback-scoped Oracle acceptance tests. Business fixtures in those tests rolled back; rejected-request security audit events use independent transactions and remain in the audit trail. The live browser used one clearly named synthetic maker and CIF.')

heading('What was exercised')
table(['Area', 'Oracle and browser evidence'], [
    ('Module 1 pages', 'Employee sign-in; Overview; Users and detail; Roles & permissions; Access requests; Audit trail; My security; account-holder key. Search, paging, conditional fields, dialogs, validation, session detail and authorization checks were sampled.'),
    ('Module 2 tabs', 'Synthetic Oracle CIF opened in Profile, KYC and evidence, Relationships and consent, and Audit. Each tab and its visible actions were inspected; keyboard arrows/Home/End, selected panel state and return to directory were verified.'),
    ('KYC guard', 'Submitting the synthetic DRAFT case without identity and address documents returned both missing-document codes and kept the case editable. Oracle tests covered upload, attestation, approval, rejection, rework, latest-version gates, expiry and closure.'),
    ('Additional Oracle checks', 'A second overlapping access grant was rejected, adjacent validity windows were accepted, and stale profile and KYC review submissions were denied after their roles were retired. Simultaneous checker decisions and the full open-browser restoration sequence remain open.'),
    ('Shared UI', 'Phone 390 px, tablet 1024 px and desktop 1440 px layouts had no document overflow; the mobile table remained horizontally scrollable. Hindi and Kannada labels were checked, then English was restored.'),
], [1.4, 5.47])

heading('Defects corrected and retested')
table(['ID', 'Observed failure', 'Correction and retest'], [
    ('D1', 'A newly linked customer could fail closure in the same Oracle clock tick.', 'Use a positive microsecond end interval; Oracle closure method passed.'),
    ('D2', 'Access target catalogue stopped after 100 users.', 'Fetch and deduplicate all pages; 101-user Oracle fixture and Node pagination checks passed.'),
    ('D3', 'BRANCH and REGION left the visible scope reference disabled; switching scope could retain stale text.', 'Synchronize JET and submitted field state. Browser recheck passed with normal focus change; disabled source value was empty.'),
    ('D4', 'Customer tabs had no keyboard traversal or tabpanel ARIA links.', 'Arrow/Home/End and roving focus added; all four Oracle customer tabs passed browser checks.'),
    ('D5', 'A late audit search could write to the next page and show a JavaScript error.', 'Check the original form/result remains connected; rapid navigation retest showed no error.'),
    ('D6', 'Treasury navigation appeared to a maker without Treasury permissions.', 'Only show it with Treasury read or liquidity permission; Oracle maker navigation no longer showed it.'),
    ('D7', 'Blank account-holder key displayed a saved success state.', 'Validate the enhanced field before saving; blank rejected, synthetic value saved and cleared in browser.'),
    ('D8', 'Configured menu label and sort order were ignored by the shell.', 'Render saved label and order for backend menu entries; source change and default Oracle menu render verified. A changed-label Oracle browser variant remains open.'),
], [.42, 2.74, 3.71])

heading('Open finding and test limits')
paragraph('The read-only Oracle preflight found one pre-existing account position mismatch: account 6 has posted balance 500000 and journal-derived balance 0. Its authoritative posting history must be reconciled before any financial closure or balance repair. No existing account, balance, journal or real customer record was changed during this run.')
paragraph('The shared Oracle schema is already bootstrapped, and no disposable schema is available. Fresh-checker creation and contaminated-signup-role tests therefore remain blocked. Parallel approval races, the complete financial closure-blocker matrix and the open-browser revocation/restoration sequence remain outstanding. A separate Oracle service stack occupied the shared run record when the controlled outage check began, so that stack was left untouched and the outage check was not run. The matrix records every case as Partial or Not run until its full sequence and variants are evidenced; no case-level acceptance signoff is claimed.')
paragraph('Final Oracle-only package verification passed without running demo-database tests. After browser testing, the synthetic CIF was confirmed CLOSED, its test identity DISABLED and its role RETIRED; only the services launched for this run were stopped. The audit trail was retained.')

heading('Case-by-case execution state')
paragraph('See MODULE-1-AND-2-DETAILED-TEST-PLAN.docx for each case’s exact inputs, steps and expected result. “Partial” means at least one related Oracle or UI check was executed, while untested variants remain. “Not run” identifies an unexecuted primary flow. The green automated counts above apply only to their 40 and 56 method inventories.')
status_counts = Counter()
for group in GROUPS:
    heading(group['title'], 2)
    rows = []
    for case in group['cases']:
        case_id = case['id']
        status = 'Not run' if case_id in NOT_RUN else 'Partial'
        status_counts[status] += 1
        evidence = NOT_RUN.get(case_id, PARTIAL_DETAILS.get(case_id, GROUP_EVIDENCE[group_key(case_id)]))
        rows.append((case_id, case['title'], status, evidence))
    table(['Case', 'Scenario', 'State', 'Evidence / open work'], rows, [.92, 2.23, .68, 3.04])

paragraph(f"Detailed-plan state: {status_counts['Partial']} Partial; {status_counts['Not run']} Not run; 0 fully signed off. The remaining cases need follow-on Oracle execution; the two fresh-schema cases cannot be completed in the current shared schema.")
paragraph('Evidence files are under Downloads/integrated/.runtime/oracle-m12/; backend acceptance source is backend/acceptance-tests/src/test/java/com/moneybags/Module12OracleTest.java. The final execution report deliberately excludes credentials, tokens, evidence content and private Oracle settings.')

OUT.parent.mkdir(parents=True, exist_ok=True)
doc.save(OUT)
assert len(CASES) == 118 and int(SUITE.attrib['tests']) == 40
assert not any(term in OUT.read_bytes() for term in (b'DB_PASSWORD', b'accessToken', b'OracleQa!2026Secure'))
print(json.dumps({'output': str(OUT), 'oracle_methods': 40, 'case_statuses': dict(status_counts), 'bytes': OUT.stat().st_size}, indent=2))
