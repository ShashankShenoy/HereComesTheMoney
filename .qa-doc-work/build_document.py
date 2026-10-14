from pathlib import Path
from collections import Counter
import hashlib, json, re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from test_cases import GROUPS

ROOT = Path(__file__).resolve().parents[1]
BASE = ROOT / 'Downloads/integrated'
OUT = BASE / 'docs/MODULE-1-AND-2-DETAILED-TEST-PLAN.docx'
WORK = Path(__file__).parent
CASES = [c for g in GROUPS for c in g['cases']]
SOURCES = {
 'S1': ('IAM pages and action handlers', ['frontend/app.js']),
 'S2': ('CIF tabs forms and action handlers', ['frontend/domains.js']),
 'S3': ('Authentication API and session transport', ['frontend/api.js', 'backend/identity-service/src/main/java/com/moneybags/iam/controller/AuthController.java', 'backend/identity-service/src/main/java/com/moneybags/iam/service/LoginService.java']),
 'S4': ('MFA rules and freshness', ['backend/identity-service/src/main/java/com/moneybags/iam/service/MfaService.java', 'backend/shared/src/main/java/com/moneybags/iam/security/MfaFreshness.java']),
 'S5': ('IAM administration contracts and rules', ['backend/shared/src/main/java/com/moneybags/iam/dto/IamDtos.java', 'backend/identity-service/src/main/java/com/moneybags/iam/controller/IamController.java', 'backend/identity-service/src/main/java/com/moneybags/iam/service/IamAdminService.java']),
 'S6': ('Scope and authority decisions', ['backend/shared/src/main/java/com/moneybags/iam/service/AccessDecisionService.java']),
 'S7': ('CIF API contracts and lifecycle', ['backend/customer-module/src/main/java/com/moneybags/cif/controller/CifController.java', 'backend/customer-module/src/main/java/com/moneybags/cif/dto/CifDtos.java', 'backend/customer-module/src/main/java/com/moneybags/cif/service/CifService.java']),
 'S8': ('Signup and customer identity key creation', ['backend/identity-service/src/main/java/com/moneybags/iam/service/CustomerSignupService.java', 'backend/identity-service/src/main/java/com/moneybags/integration/CustomerHashController.java']),
 'S9': ('Financial audit UI integration', ['frontend/financial-audit.test.js', 'docs/FINANCIAL-AUDIT-TESTING.txt']),
 'S10': ('Account holder key UI and backend', ['frontend/banking.js', 'backend/shared/src/main/java/com/moneybags/integration/CustomerHashService.java']),
 'S11': ('Evidence storage validation and review', ['backend/customer-module/src/main/java/com/moneybags/cif/service/KycDocumentService.java', 'backend/shared/src/main/java/com/moneybags/common/security/PrivateData.java']),
 'S12': ('CIF closure response in IAM', ['backend/shared/src/main/java/com/moneybags/iam/service/CustomerAccessLifecycle.java']),
 'S13': ('Shared controls and presentation', ['frontend/ui.js', 'frontend/jet.js', 'frontend/i18n.js', 'frontend/design.js']),
 'S14': ('Local setup and implementation boundaries', ['README.md', 'docs/FEATURE-COVERAGE.md', 'docs/MICROSERVICES.md']),
}

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)
sec.top_margin = sec.bottom_margin = Inches(.65)
sec.left_margin = sec.right_margin = Inches(.7)
sec.header_distance = sec.footer_distance = Inches(.3)
for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3'):
    st = doc.styles[name]
    st.font.name = 'Arial'
    st.font.color.rgb = RGBColor(0,0,0)
    st.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Arial')
    for node in st.element.xpath('.//w:pBdr'):
        node.getparent().remove(node)
    for node in st.element.xpath('.//w:color'):
        for attr in ('themeColor','themeTint','themeShade'):
            node.attrib.pop(qn('w:'+attr), None)
normal = doc.styles['Normal']
normal.font.size = Pt(10)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.04
for name, size, before, after in [('Title',28,0,12),('Subtitle',13,0,9),('Heading 1',18,16,8),('Heading 2',14,12,7),('Heading 3',11,10,4)]:
    st=doc.styles[name]; st.font.size=Pt(size)
    st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after)
    st.paragraph_format.keep_with_next=True
doc.styles['Title'].font.bold=True
doc.core_properties.title='Moneybags Module 1 and Module 2 Detailed Test Plan'
doc.core_properties.subject='IAM and Customers and KYC functional interaction and integration test specification'
doc.core_properties.author='Moneybags QA'
doc.core_properties.keywords='M01 IAM M02 CIF KYC test cases regression UAT'

def p(text='', bold=False, style=None):
    para=doc.add_paragraph(style=style)
    run=para.add_run(text); run.bold=bold
    return para

def lead(label,text):
    para=doc.add_paragraph()
    para.add_run(label+' ').bold=True
    para.add_run(text)
    return para

def h(text,level=1):
    return doc.add_heading(text,level)

def table(headers, rows, widths):
    t=doc.add_table(rows=1, cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for i,w in enumerate(widths):t.columns[i].width=Inches(w)
    for i,x in enumerate(headers):t.rows[0].cells[i].text=x
    for row in rows:
        cells=t.add_row().cells
        for i,x in enumerate(row):cells[i].text=str(x)
    for ri,row in enumerate(t.rows):
        trPr=row._tr.get_or_add_trPr()
        cant=OxmlElement('w:cantSplit');trPr.append(cant)
        if ri==0:
            repeat=OxmlElement('w:tblHeader');trPr.append(repeat)
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            pr=cell._tc.get_or_add_tcPr()
            sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE6EF' if ri==0 else ('F5F7F9' if ri%2==0 else 'FFFFFF'));pr.append(sh)
            margins=OxmlElement('w:tcMar')
            for edge,value in [('top',75),('bottom',75),('left',90),('right',90)]:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),str(value));e.set(qn('w:type'),'dxa');margins.append(e)
            pr.append(margins)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(1)
                para.paragraph_format.line_spacing=1.02
                if ri==0: para.paragraph_format.keep_with_next=True
                for run in para.runs:run.font.size=Pt(9);run.bold=ri==0
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
    t._tbl.tblPr.append(borders)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)
    return t

def field(para,instruction):
    f=OxmlElement('w:fldSimple');f.set(qn('w:instr'),instruction)
    run=OxmlElement('w:r');text=OxmlElement('w:t');text.text='1';run.append(text);f.append(run);para._p.append(f)

footer=sec.footer.paragraphs[0]
footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer.add_run('Moneybags QA  |  Page ').font.size=Pt(8)
field(footer,'PAGE');footer.add_run(' of ').font.size=Pt(8);field(footer,'NUMPAGES')

p('Moneybags Module 1 and Module 2 Detailed Test Plan',style='Title')
p('Identity and Access Management and Customers and KYC',style='Subtitle')
p('Functional coverage and execution specification',bold=True)
p('Version 1.0  |  14 October 2026')
p(f'This document defines {len(CASES)} traceable test cases for QA and acceptance testers. Execute the cases to verify every in-scope page, customer record tab, dialog, state transition and cross-module handoff. Each case includes setup, concrete actions and an expected outcome so that failures can be reproduced and reviewed.')
lead('Execution status', 'All cases start as Not run. This is a test specification and execution worksheet, not a claim that the application has passed these scenarios. Record the tested build, actual results and evidence before assigning Pass or Fail.')
table(['Coverage area','Cases','Primary concern'],[
    ['Module 1',sum(len(g['cases']) for g in GROUPS if g['code'].startswith('M1')),'Identity, access approval, security, navigation and audit'],
    ['Module 2',sum(len(g['cases']) for g in GROUPS if g['code'].startswith('M2')),'Customer information, KYC, evidence, consent and lifecycle'],
    ['Shared and integrated flows',sum(len(g['cases']) for g in GROUPS if g['code'].startswith('X')),'Handoffs, revocation, failure recovery and common interactions'],
], [1.55,.65,4.67])
sp=p('Scope and use',bold=True);sp.runs[0].font.size=Pt(14);sp.paragraph_format.keep_with_next=True
p('Module 1 includes the six Administration pages: Overview, Users, Roles & permissions, Access requests, Audit trail and My security. Sign in, customer registration and Account-holder access key are included because they create or control the same identity and access state. Overview here is the IAM page, distinct from the general banking overview.')
p('Module 2 includes Customer directory and all four customer tabs: Profile, KYC and evidence, Relationships and consent, and Audit. Case evidence is a panel within KYC and evidence. IAM detail sections are panels or dialogs, not additional tabs.')
p('Account, deposit, payment, product, card and loan functionality is checked only where it consumes customer eligibility or prevents CIF closure. Full functional testing of those modules is outside this document. Use synthetic customers and isolated fixtures for closures, lockouts, concurrency and injected failures.')

doc.add_page_break()
cp=p('Contents',bold=True);cp.runs[0].font.size=Pt(18);cp.paragraph_format.keep_with_next=True
p('Use the entries below or Word Navigation to reach each area. Individual cases are grouped by their visible screen and workflow.')
tocp=doc.add_paragraph()
toc=OxmlElement('w:fldSimple');toc.set(qn('w:instr'),'TOC \\o "1-2" \\h \\z \\u')
tocp._p.append(toc)

doc.add_page_break()
h('Test environment and fixtures')
p('Record the frontend and backend build identifiers, database profile, browser version, active feature flags, bank business timezone and tester before starting. The source snapshot for this specification is identified in the source register. Run the normal KYC suite with demo signup disabled; test demo registration separately with the flag enabled.')
table(['Fixture','Required setup','Use'],[
 ['A and B','Two separate authorized administrators; B is independent checker','Provision roles and approve access without self approval'],
 ['E1 maker','ACTIVE EMPLOYEE; B1 CIF_READ, CIF_CREATE, CIF_UPDATE, KYC_UPLOAD','Create and propose customer data; submit cases'],
 ['E2 reviewer','ACTIVE EMPLOYEE; B1 CIF_READ and KYC_REVIEW; distinct from E1','Assign review, inspect evidence and decide KYC'],
 ['E3 and E4','Read-only B1 employee and B2-only employee; no hidden global assignment','Least privilege and branch isolation'],
 ['C1 and C2','CUSTOMER identities with different ACTIVE CIF links and SELF grants','Ownership and cross-customer denial'],
 ['No-access identity','ACTIVE login with no module permissions','Direct API authorization negative tests'],
 ['Customer state set','PROSPECT, PENDING_VERIFICATION, ACTIVE, RESTRICTED, DORMANT, CLOSED; KYC states current/due/rejected','Status and effective eligibility tests'],
 ['KYC state set','DRAFT, DOCUMENTS_PENDING, MORE_INFO_REQUIRED, SUBMITTED, UNDER_REVIEW, APPROVED, REJECTED, EXPIRED','UI controls and API state guards'],
 ['Evidence set','Synthetic PDF, PNG and JPEG; versions; clean/pending/quarantined/failed fixtures; size boundaries','Upload, review, download and integrity checks'],
 ['Dependency set','Open primary/joint account, active term deposit, live/pending card, loan application/facility; zero-obligation CIF','Closure blockers and successful closure'],
], [1.2,3.4,2.27])
lead('Local setup', 'The repository start.ps1 builds and starts the local stack; the root README describes http://localhost:5173 and /employee/login. The local stack resets synthetic data when stopped. Use a dedicated test run and the environment credential handoff; never embed passwords, holder keys, MFA secrets or tokens in evidence.')
lead('Clock and configuration', 'Use a controlled clock or purpose-built fixtures for expiry boundaries. Do not change the workstation clock. Confirm actual configuration; defaults in LoginService are 5 failed attempts, 15-minute credential lock, 30-minute idle session and 8-hour absolute session. Fresh MFA is accepted through age 300 seconds inclusive.')
lead('Inspection', 'Use UI plus an authorized API client and read-only database inspection. All endpoint paths in this document are relative to /api/v1. Fault injection and database immutability probes use disposable fixtures. Capture before/after counts, statuses and versions; restore only the fixtures changed by the test.')
lead('Data naming', 'Use a run prefix such as QA_M12_20261014 and unique usernames/identifier values per accepted variant. Use qa.customer@example.test and synthetic telephone numbers. Take dates relative to the application business date, and never use real identity documents.')

h('Execution method and acceptance criteria')
p('Each numbered case can contain several parameter variants. Run every stated variant, assign a suffix such as a, b or c, and record an individual result. A case passes only when all required variants pass. A missing prerequisite is Blocked, not Pass. A deliberately unavailable deployment feature is Not applicable only with a recorded reason and reviewer acceptance.')
table(['Priority','Meaning','Release decision'],[
 ['P0','Authorization, segregation of duties, privacy, atomic lifecycle or core onboarding','All in-scope P0 variants pass; no unresolved critical defect'],
 ['P1','Functional correctness, field rules, navigation and recovery','All required P1 variants pass or an explicit accepted exception is recorded'],
 ['P2','Large-list usability, responsive layout and localization depth','Record findings and disposition; inaccessible core actions remain release relevant'],
], [0.7,3.05,3.12])
p('For a successful mutation, verify the visible confirmation, persisted result after reload, correct actor/resource audit and relevant downstream effect. For a rejected mutation, verify the error and unchanged business data. Do not require every failure to produce a domain event: check the applicable denial audit without treating a rejected action as a committed change.')
p('Attach redacted evidence to the case ID: browser screenshot, request method/path, status and response code, correlation ID, before/after row versions, and relevant audit/event references. Record timing for race/expiry tests. Avoid raw credentials or unmasked sensitive fields in screenshots, URLs, logs and exported traces.')

h('Page and action coverage inventory')
p('The ranges below map the visible navigation and every action family to detailed cases. Shared dialog, authorization, keyboard and responsive checks apply to every row. The case catalogue follows the same order.')
coverage=[]
for g in GROUPS:
    coverage.append([g['location'],g['title'],f"{g['cases'][0]['id']} to {g['cases'][-1]['id']}"])
table(['Page or tab','Interaction coverage','Test cases'],coverage,[2.25,2.85,1.77])
table(['Module 2 tab','Panels and controls to verify'],[
 ['Profile','Customer profile and Additional record fields; Propose profile change; Change status; Names & history; Contacts and Add contact; Identifiers and Add identifier; Addresses and Add address.'],
 ['KYC and evidence','KYC cases; Open new case; Open evidence; Upload document; Download; Review evidence; Submit case; Assign officer; Decide KYC; document versions and review history.'],
 ['Relationships and consent','Consent evidence; Capture consent; Withdraw; Party relationships; Add relationship; End on eligible outgoing relations.'],
 ['Audit','Customer audit event, result, actor and time; scope, reload and customer switching.'],
], [1.55,5.32])

h('Workflow and state expectations')
table(['Flow','Required sequence and checkpoints'],[
 ['Staff access','Create identity > Activate > Request grant > Independent checker approves > Entitlement effective. Pending or rejected request gives no new access.'],
 ['Normal customer onboarding','Create CIF and DRAFT case > Propose contact, identifier and address > Upload IDENTITY and ADDRESS > Submit > Optionally assign officer > Independent evidence attestation > Independent KYC approval > ACTIVE and VERIFIED for initial onboarding.'],
 ['KYC rework','SUBMITTED or UNDER_REVIEW > REQUEST_INFO > MORE_INFO_REQUIRED > Original maker edits and uploads > Resubmit > Reattest latest documents > Decide.'],
 ['Periodic review','Open new case on verified CIF > KYC PENDING and ACTIVE becomes RESTRICTED > Propose revisions > Independent approval > KYC VERIFIED with restriction preserved > Separate status review to ACTIVE.'],
 ['Digital customer access','Create and activate customer login > Link existing ACTIVE CIF > Governed SELF assignment > Sign in > Access only linked authorized data.'],
 ['CIF closure','Resolve financial blockers > Close CIF > Expire open cases and reject pending facts > Withdraw consents and end relationships > Revoke links and sessions > Disable login only if no other active CIF remains.'],
 ['Demo signup','Feature flag enabled > Adult signup > ACTIVE identity and assumed VERIFIED CIF plus SELF assignment. Test separately from normal evidence-based KYC.'],
], [1.4,5.47])
table(['Identity current state','Allowed next states through Change status'],[
 ['PENDING_ACTIVATION','ACTIVE, DISABLED, CLOSED'],['ACTIVE','AUTH_LOCKED, SUSPENDED, DISABLED, CLOSED'],
 ['AUTH_LOCKED','ACTIVE, SUSPENDED, DISABLED, CLOSED'],['SUSPENDED','ACTIVE, DISABLED, CLOSED'],
 ['DISABLED','ACTIVE, CLOSED'],['CLOSED','None'],
], [2.0,4.87])
table(['KYC state family','Allowed work and guards'],[
 ['DRAFT / DOCUMENTS_PENDING / MORE_INFO_REQUIRED','Original maker can propose facts, upload and submit, with matching permission. DOCUMENTS_PENDING is a supported state fixture; there is no separate UI command that must create it.'],
 ['SUBMITTED / UNDER_REVIEW','Maker editing is frozen. Independent eligible reviewer can inspect/attest; if assigned, only assigned officer decides. Assignment sets UNDER_REVIEW.'],
 ['APPROVED / REJECTED / EXPIRED','Terminal case evidence history remains; upload, submit and review are not allowed. A nonclosed CIF may open another case.'],
], [2.35,4.52])

current_module=None
for g in GROUPS:
    module='Module 1 detailed test cases' if g['code'].startswith('M1') else 'Module 2 detailed test cases' if g['code'].startswith('M2') else 'Integrated and shared test cases'
    if module!=current_module:
        if current_module is not None:doc.add_page_break()
        h(module);current_module=module
    h(g['title'],2)
    loc=lead('Location',g['location']+'. Source references: '+g['source']+'.')
    loc.paragraph_format.keep_with_next=True
    for c in g['cases']:
        # Case IDs remain literal in the body for search and evidence mapping.
        hp=h(c['title'],3)
        meta=p(f"{c['id']}  |  {c['priority']}  |  {c['mode']}  |  Status Not run")
        meta.paragraph_format.space_after=Pt(3)
        meta.runs[0].font.size=Pt(9);meta.runs[0].bold=True
        paras=[hp,meta,lead('Preconditions',c['pre'])]
        for i,step in enumerate(c['steps'],1):
            sp=p(f'{i}. {step}');sp.paragraph_format.space_after=Pt(2);paras.append(sp)
        paras.append(lead('Expected result',c['expected']))
        rec=p('Actual result and variant IDs: ____________________   Evidence or defect: ____________________')
        rec.runs[0].font.size=Pt(8);rec.paragraph_format.space_after=Pt(6)
        paras.append(rec)
        # Keep each short case together to avoid detached steps or result lines.
        for para in paras[:-1]:para.paragraph_format.keep_with_next=True

doc.add_page_break()
h('Field validation and boundary matrix')
p('For each maximum length N, test blank where applicable, N minus 1, N and N plus 1. Use API requests to test values the browser cannot submit. Record UI and API differences separately. Whitespace-only is invalid for required strings. Dates use application time; validity end must be strictly after start unless the case specifies another rule.')
fields=[
 ['Username','3 to 120; A-Z a-z 0-9 . _ @ -; case-insensitive uniqueness','M1-AUTH-08; M1-USR-04'],
 ['New password','At least 10 characters; at most 72 UTF-8 bytes; test multibyte values','M1-USR-04; M1-SEC-02'],
 ['Employee reference','Required only for EMPLOYEE; max 80','M1-USR-02,04'],
 ['IAM reason','Required; max 500 including status, reset and access requests','M1-USR-06,08; M1-ACC-02'],
 ['Role code and display','Code 3 to 80, uppercase letter first, letters/digits/underscore; display max 120','M1-ROLE-02,03'],
 ['Role permissions','Non-null list; maximum 100 IDs; unique known IDs; empty list is permitted by DTO','M1-ROLE-03'],
 ['Permission code','3 to 100 uppercase code; IAM_ and SYSTEM_ reserved','M1-ROLE-05'],
 ['Permission resource and action','Required, each max 60; step-up is boolean','M1-ROLE-04,05'],
 ['Access scope and dimensions','Scope ref max 80; branch/region required; global/self absent; at most 30 dimension records; value max 80','M1-ACC-01,02,06'],
 ['Access dates','validFrom required; validTo optional, after from and not already ended; grant effective at from, ceases at to','M1-ACC-03,07,08'],
 ['Authority code and currency','Code 3 to 80 uppercase code; currency exactly three uppercase letters or *','M1-ROLE-06,07'],
 ['Authority amount and rate','Nonnegative; amount max 16 integer and 4 fraction digits; rate max 4 integer and 8 fraction digits; at least one limit','M1-ROLE-06,07'],
 ['Menu fields','Label max 120; route max 240 with leading slash and supported path characters; sort 0 through 99999','M1-MENU-01,02'],
 ['MFA code','Exactly six decimal digits; include leading zero; 30-second TOTP with previous/current/next slot acceptance and replay protection','M1-AUTH-05; M1-SEC-03,05'],
 ['CIF legal name','Required max 300; Unicode and harmless markup-like text; no assumed name uniqueness','M2-DIR-04,06; M2-PRO-02'],
 ['Party dates','INDIVIDUAL requires DOB only; ORGANIZATION requires incorporation only; neither future; normal CIF flow has no public-signup adult gate','M2-DIR-04,05'],
 ['Home branch and segment','Branch required max 80, UI trims; segment optional max 40; pending segment effective only after approval','M2-DIR-06; M2-PRO-02,03'],
 ['CIF revisions and reasons','rowVersion >= 1; profile/status/submit reason max 1000 in API; shared UI Reason limits to 500','M2-PRO-03; M2-KYC-04; M2-LIFE-01'],
 ['Contacts','MOBILE or EMAIL; value required max 240; email/mobile implementation regex applied','M2-PRO-05,06'],
 ['Identifiers','Type required max 40; value max 200; issuer optional max 80; expiry optional but strictly after today','M2-PRO-07,08'],
 ['Address text','line1 required max 240; line2 optional max 240; city required max 120; region optional max 120; postal optional max 24','M2-PRO-09,10'],
 ['Address country and evidence','Country required exactly two uppercase letters; optional evidence ID max 36 with case/CIF/type/maker ownership','M2-PRO-09,10'],
 ['Case type and assignment','ONBOARDING/PERIODIC/REMEDIATION; officer ID required max 36; valid current rowVersion','M2-KYC-01,05,06'],
 ['KYC decision','APPROVE/REJECT/REQUEST_INFO; reasonCode required max 80; comments required max 2000; approval risk and future date required','M2-KYC-07,08,10,11'],
 ['Document upload','Actual PDF/PNG/JPEG signatures; >0 and <=5242880 bytes; type uppercase code max 40; UI choices IDENTITY/ADDRESS/OTHER','M2-DOC-01,02,03'],
 ['Evidence review','CLEAN/QUARANTINED/FAILED; VERIFIED/REJECTED; nonclean cannot be VERIFIED; reason max 1000 API and 500 UI','M2-DOC-04,05'],
 ['Consent','Purpose required max 80; channel max 40; evidence reference max 240; one active grant per CIF/purpose','M2-REL-01,02'],
 ['Relationship','Target CIF required max 36; type required max 32; NONE/VIEW/TRANSACT/REPRESENT; optional end strictly future','M2-REL-03,04,05'],
 ['Record identifiers','Unknown/malformed/other-customer IDs; never substitute a different existing record; direct endpoint checks','M1-USR-10,11; X-UI-06'],
]
table(['Field family','Boundary and rule','Coverage'],fields,[1.45,3.75,1.67])

h('Known implementation distinctions to verify')
p('These checks separate specified behavior from product improvements. Log observed discrepancies as defects or requirement decisions; do not silently alter expected results to make a run pass.')
table(['Area','Current distinction and test treatment'],[
 ['Demo registration','Signup and admin demo-customer creation deliberately assume KYC; a passing demo path does not replace the full evidence-based onboarding flow.'],
 ['Evidence scan','Review evidence is an explicit checker attestation. Signature checks identify formats but do not prove malware scanning or complete file-content validation.'],
 ['List sizes','Users UI pages 15 rows; role/access catalog loads first 100 users; access queue caps at 200; CIF directory caps at 200 and has no UI pagination. Test target discoverability.'],
 ['Navigation metadata','Backend stores labels/order, but shell uses fixed page titles/navigation order and explicitly adds certain pages. Test visible behavior and record configurable-menu gaps.'],
 ['CIF and account authority','Digital CIF link, party relationship and account operating mandate are distinct. TRANSACT on a party relationship must not be treated as a completed account authorization workflow.'],
 ['Mutation replay','The browser sends Idempotency-Key on mutations, but that does not prove every IAM/CIF handler enforces replay. Verify duplicate outcomes per operation and rely only on documented state/version/uniqueness protections.'],
 ['Evidence metadata','CIF record detail omits sensitive encryption/lookup fields and storage paths. Case metadata currently includes document SHA256_HEX for integrity; do not confuse it with a bearer token or plaintext identifier.'],
 ['Deployment validation','Local in-memory and Oracle deployments require separate recorded runs for locking, timestamp precision, constraints, encrypted storage and audit immutability. Historical repository reports are reference material, not execution evidence for this plan.'],
], [1.5,5.37])

h('Test run record and completion review')
table(['Run field','Recorded value'],[
 ['Run ID and date','____________________________________________'],
 ['Build or release identifier','____________________________________________'],
 ['Database and service profile','____________________________________________'],
 ['Browser version and viewport','____________________________________________'],
 ['Feature flags and clock timezone','____________________________________________'],
 ['Tester and reviewer','____________________________________________'],
 ['Passed / Failed / Blocked / Not run / Not applicable','____________________________________________'],
 ['Defects and accepted exceptions','____________________________________________'],
], [2.55,4.32])
p('Create an execution record per case variant with: case ID, input variant, build, tester/time, actual result, result status, evidence link, defect ID and retest result. Use the line beneath each case for concise notes and link larger evidence from the test system. Preserve the first failed result when a retest passes.')
table(['Coverage signoff','Required evidence','Result'],[
 ['Module 1 pages','All six Administration pages, entry flows and holder-key page executed','Not run'],
 ['Module 2 Profile','All five panels plus profile/status dialogs executed','Not run'],
 ['Module 2 KYC and evidence','Every case state, decision, upload/review/download and rework flow executed','Not run'],
 ['Module 2 Relationships and consent','Create, withdraw/end, duplicates and authority isolation executed','Not run'],
 ['Module 2 Audit','Correct resource timeline and scope checks executed','Not run'],
 ['Integrated flows','Normal onboarding, linking, revocation, expiry, closure and outages executed','Not run'],
 ['Shared interactions','Every dialog cancel/save/error path and supported accessibility/layout checks executed','Not run'],
], [1.9,3.9,1.07])
p('Exit when required variants are executed, all P0 cases pass, blocking defects are resolved, and any remaining exception has an owner and explicit disposition. Re-run affected cases and the connected end-to-end flow after fixes. A green automated suite alone does not sign off manual tab, role and evidence handling coverage.')

h('Source register')
p('Source references resolve from Downloads/integrated in the supplied project. They identify the inspected implementation for this specification. The full SHA-256 snapshot is retained with the document build data; short fingerprints below identify the exact files. No historical test result is claimed as a new run.')
manifest=[]
for sid,(title,paths) in SOURCES.items():
    sp=p(sid+' '+title,bold=True);sp.runs[0].font.size=Pt(11);sp.paragraph_format.keep_with_next=True
    for rel in paths:
        src=BASE/rel
        digest=hashlib.sha256(src.read_bytes()).hexdigest()
        manifest.append({'source':sid,'path':rel,'sha256':digest})
        para=p(rel+'  |  SHA256 '+digest[:16])
        para.paragraph_format.space_after=Pt(3)
        for run in para.runs:run.font.size=Pt(8)

OUT.parent.mkdir(parents=True,exist_ok=True)
doc.save(OUT)
(WORK/'source_manifest.json').write_text(json.dumps(manifest,indent=2),encoding='utf-8')
(WORK/'case_inventory.json').write_text(json.dumps(GROUPS,indent=2),encoding='utf-8')
ids=[c['id'] for c in CASES]
assert len(ids)==len(set(ids))
assert all(len(c['steps'])>=2 and len(c['expected'])>70 for c in CASES)
assert all((BASE/rel).is_file() for _,paths in SOURCES.values() for rel in paths)
print(json.dumps({'output':str(OUT),'cases':len(CASES),'groups':len(GROUPS),'priorities':dict(Counter(c['priority'] for c in CASES)),'bytes':OUT.stat().st_size},indent=2))
